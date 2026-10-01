"""Cube ACR transcripts from the public Drive folder via Drive API v3.

The anonymous folder HTML only lists the first ~50 subfolders and never
matched *-transcript.docx, so ingest stayed at 0. This module pages the
Drive API, walks day folders, and downloads .docx / Google Docs.
"""

from __future__ import annotations

import logging
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Iterable

import requests

from crmbrain.config import (
    CDT,
    Settings,
    date_window_cdt,
    digits_phone,
    is_personal,
    lookback_dates_cdt,
    now_utc,
    settings_lookback_start,
    today_and_yesterday_cdt,
)
from crmbrain.google_auth import (
    DRIVE_READONLY,
    granted_scopes,
    has_drive_scope,
    oauth_configured,
    refresh_access_token,
)
from crmbrain.models import Engagement
from crmbrain.policy import looks_like_html

logger = logging.getLogger(__name__)

DRIVE_FILES = "https://www.googleapis.com/drive/v3/files"
FOLDER_MIME = "application/vnd.google-apps.folder"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GDOC_MIME = "application/vnd.google-apps.document"
DOCX_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
DATE_FOLDER_RE = re.compile(r"^(20\d{2}-\d{2}-\d{2})$")
TITLE_DATE_RE = re.compile(r"(20\d{2}-\d{2}-\d{2})")
TITLE_HMS_RE = re.compile(r"\b(\d{1,2}[-:]\d{2}[-:]\d{2})\b")
TITLE_HM_RE = re.compile(r"\b(\d{1,2}[:]\d{2})\b")
PAREN_INTL_PHONE_RE = re.compile(
    r"\(\s*\+?1\s*(\d{3})\s*[-.\s]?\s*(\d{3})\s*[-.\s]?\s*(\d{4})\s*\)"
)
PAREN_US_PHONE_RE = re.compile(r"\(\s*(\d{3})\s*\)\s*(\d{3})\s*[-.\s]?\s*(\d{4})")
NAKED_PHONE_RE = re.compile(r"(?:\+?1[\s\-.]?)?(?:\(?\d{3}\)?[\s\-.]?\d{3}[\s\-.]?\d{4})")
ARROW_RE = re.compile(r"[↗↘↑↓➔➜]+")
PHONE_WORD_RE = re.compile(r"\(\s*phone\s*\)", re.I)
TRANSCRIPT_TAIL_RE = re.compile(r"(?i)\s*-\s*transcript(?:\.(?:docx|doc|gdoc|txt))?$")
EXT_TAIL_RE = re.compile(r"(?i)\.(?:docx|doc|gdoc|txt)$")
KIND_RANK = {
    "docx_transcript": 0,
    "docx": 1,
    "gdoc": 2,
    "text": 3,
}


class CubeAuthError(Exception):
    """Missing drive.readonly or GOOGLE_API_KEY — stale-source warning, not an error."""


@dataclass(frozen=True)
class DriveAuth:
    mode: str
    headers: dict[str, str]
    key: str = ""


@dataclass
class DriveFile:
    file_id: str
    name: str
    mime_type: str
    modified_time: datetime | None = None
    md5: str = ""
    folder_date: str = ""


def _parse_drive_time(raw: str) -> datetime | None:
    if not raw:
        return None
    text = str(raw).replace("Z", "+00:00")
    try:
        stamp = datetime.fromisoformat(text)
    except ValueError:
        return None
    return stamp if stamp.tzinfo else stamp.replace(tzinfo=timezone.utc)


def cube_window_start(settings: Settings, *, backfill: bool = False) -> datetime:
    """Cycle lookback, or the full CUBE_LOOKBACK_DAYS window on first backfill."""
    cycle = settings_lookback_start(settings)
    if not backfill:
        return cycle
    days = int(getattr(settings, "cube_lookback_days", 14) or 14)
    return min(cycle, now_utc() - timedelta(days=max(1, days)))


def cube_listing_dates(settings: Settings, *, backfill: bool = False) -> list[str]:
    if backfill:
        return date_window_cdt(int(getattr(settings, "cube_lookback_days", 14) or 14))
    return lookback_dates_cdt(settings) or today_and_yesterday_cdt()


def resolve_drive_auth(settings: Settings) -> DriveAuth:
    if oauth_configured(settings):
        try:
            scopes = granted_scopes(settings)
            if has_drive_scope(scopes):
                token = refresh_access_token(settings)
                return DriveAuth(
                    mode="oauth",
                    headers={"Authorization": f"Bearer {token}"},
                )
        except Exception as exc:
            logger.warning("cube drive oauth failed: %s", exc)
    key = (settings.google_api_key or "").strip()
    if key:
        return DriveAuth(mode="api_key", headers={}, key=key)
    if oauth_configured(settings):
        raise CubeAuthError(
            "Gmail OAuth token is missing drive.readonly; "
            "re-consent that scope or set GOOGLE_API_KEY for the public Cube folder"
        )
    raise CubeAuthError("drive.readonly scope or GOOGLE_API_KEY missing")


def drive_get(auth: DriveAuth, url: str, params: dict[str, Any] | None = None) -> Any:
    merged = dict(params or {})
    if auth.key:
        merged.setdefault("key", auth.key)
    resp = requests.get(url, headers=auth.headers, params=merged, timeout=45)
    if resp.status_code in {401, 403}:
        raise PermissionError(f"drive api {resp.status_code}")
    resp.raise_for_status()
    if "application/json" in (resp.headers.get("content-type") or "") or url.rstrip("/").endswith("/files"):
        if "alt" in merged and merged.get("alt") == "media":
            return resp.content
        try:
            return resp.json()
        except ValueError:
            return resp.content
    return resp.content


def list_drive_children(
    auth: DriveAuth,
    folder_id: str,
    *,
    q_extra: str = "",
    order_by: str = "modifiedTime desc",
    page_size: int = 100,
    modified_after: datetime | None = None,
) -> list[dict]:
    """Page Drive API v3 children. page_size is capped by the API, not by HTML's 50."""
    files: list[dict] = []
    page_token = ""
    extra = f" {q_extra}" if q_extra else ""
    if modified_after is not None:
        extra += f" and modifiedTime > '{modified_after.astimezone(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')}'"
    query = f"'{folder_id}' in parents and trashed = false{extra}"
    while True:
        params: dict[str, Any] = {
            "q": query,
            "fields": "nextPageToken, files(id, name, mimeType, modifiedTime, md5Checksum, size)",
            "orderBy": order_by,
            "pageSize": page_size,
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
        if page_token:
            params["pageToken"] = page_token
        data = drive_get(auth, DRIVE_FILES, params)
        if not isinstance(data, dict):
            break
        files.extend(data.get("files") or [])
        page_token = data.get("nextPageToken") or ""
        if not page_token:
            break
    return files


def _as_drive_file(row: dict, folder_date: str = "") -> DriveFile:
    return DriveFile(
        file_id=str(row.get("id") or ""),
        name=str(row.get("name") or ""),
        mime_type=str(row.get("mimeType") or ""),
        modified_time=_parse_drive_time(str(row.get("modifiedTime") or "")),
        md5=str(row.get("md5Checksum") or ""),
        folder_date=folder_date,
    )


def _in_window(stamp: datetime | None, start: datetime, folder_date: str, allowed_dates: set[str]) -> bool:
    if folder_date and folder_date in allowed_dates:
        return True
    if stamp is None:
        return bool(folder_date) and folder_date in allowed_dates
    if stamp.tzinfo is None:
        stamp = stamp.replace(tzinfo=timezone.utc)
    return stamp >= start


def list_transcript_candidates(
    settings: Settings,
    auth: DriveAuth | None = None,
    *,
    dates: Iterable[str] | None = None,
    page_size: int = 100,
    backfill: bool = False,
) -> list[DriveFile]:
    """Recursive day-folder listing, newest first, lookback-bounded."""
    auth = auth or resolve_drive_auth(settings)
    start = cube_window_start(settings, backfill=backfill)
    allowed = set(dates or cube_listing_dates(settings, backfill=backfill))
    if dates:
        allowed = set(dates)
        try:
            earliest = min(datetime.fromisoformat(d).replace(tzinfo=timezone.utc) for d in allowed)
            start = min(start, earliest)
        except ValueError:
            pass
    root = list_drive_children(auth, settings.cube_folder, page_size=page_size)
    folders: list[DriveFile] = []
    files: list[DriveFile] = []
    for row in root:
        item = _as_drive_file(row)
        if not item.file_id:
            continue
        if item.mime_type == FOLDER_MIME:
            m = DATE_FOLDER_RE.match(item.name.strip())
            item.folder_date = m.group(1) if m else ""
            if _in_window(item.modified_time, start, item.folder_date, allowed) or (
                item.folder_date and item.folder_date in allowed
            ):
                folders.append(item)
            continue
        if _in_window(item.modified_time, start, "", allowed):
            files.append(item)
    for folder in folders:
        children = list_drive_children(auth, folder.file_id, page_size=page_size)
        for row in children:
            item = _as_drive_file(row, folder_date=folder.folder_date)
            if item.mime_type == FOLDER_MIME or not item.file_id:
                continue
            if _in_window(item.modified_time, start, folder.folder_date, allowed):
                files.append(item)
    files.sort(key=lambda f: f.modified_time or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
    return files


def to_e164(raw: str) -> str:
    """US numbers as +1XXXXXXXXXX. Empty when there are not enough digits."""
    digits = digits_phone(raw)
    if len(digits) == 10:
        return f"+1{digits}"
    if len(digits) >= 11 and digits.startswith("1"):
        return f"+{digits[:11]}"
    return ""


def _normalize_hms(raw: str) -> str:
    parts = re.split(r"[-:.]", raw)
    try:
        nums = [int(p) for p in parts if p != ""]
    except ValueError:
        return ""
    if len(nums) == 2:
        nums.append(0)
    if len(nums) != 3:
        return ""
    hh, mm, ss = nums
    return f"{hh:02d}:{mm:02d}:{ss:02d}"


def parse_cube_title(name: str) -> dict[str, str]:
    """Name, E.164 phone, date, and Chicago-local time from a real Cube title."""
    rest = name or ""
    date = ""
    dm = TITLE_DATE_RE.search(rest)
    if dm:
        date = dm.group(1)
        rest = rest[: dm.start()] + " " + rest[dm.end() :]

    time_s = ""
    tm = TITLE_HMS_RE.search(rest) or TITLE_HM_RE.search(rest)
    if tm:
        time_s = _normalize_hms(tm.group(1))
        rest = rest[: tm.start()] + " " + rest[tm.end() :]

    phone = ""
    for pattern in (PAREN_INTL_PHONE_RE, PAREN_US_PHONE_RE, NAKED_PHONE_RE):
        pm = pattern.search(rest)
        if not pm:
            continue
        phone = to_e164(pm.group(0))
        rest = rest[: pm.start()] + " " + rest[pm.end() :]
        break

    rest = ARROW_RE.sub(" ", rest)
    rest = PHONE_WORD_RE.sub(" ", rest)
    rest = TRANSCRIPT_TAIL_RE.sub(" ", rest)
    rest = EXT_TAIL_RE.sub(" ", rest)
    rest = re.sub(r"(?i)\btranscript\b", " ", rest)
    rest = re.sub(r"[()[\]{}<>_]+", " ", rest)
    rest = re.sub(r"[\-–—.,:;!]+", " ", rest)
    person = " ".join(rest.split()).strip()
    if person and not re.search(r"[A-Za-z]", person):
        person = ""
    return {"date": date, "time": time_s, "phone": phone, "name": person}


def cube_call_key(meta: dict[str, str]) -> str:
    """Same call from two accounts: last-10 phone (or name) + date + time."""
    digits = digits_phone(meta.get("phone") or "")
    ident = digits[-10:] if len(digits) >= 10 else " ".join((meta.get("name") or "").lower().split())
    return f"{ident}|{meta.get('date') or ''}|{meta.get('time') or ''}"


def _kind_rank(kind: str, filename: str) -> int:
    low = (filename or "").lower()
    if "transcript" in low and (kind in {"docx_transcript", "docx"} or low.endswith(".docx")):
        return 0
    return KIND_RANK.get(kind, 9)


def should_skip_cube_call(name: str = "", phone: str = "", title: str = "") -> bool:
    return is_personal(name=name or title, phone=phone)


def file_kind(window: str) -> str:
    w = (window or "").lower()
    if "transcript.docx" in w or (".docx" in w and "transcript" in w):
        return "docx_transcript"
    if ".docx" in w:
        return "docx"
    if "google-apps.document" in w or w.endswith(".gdoc"):
        return "gdoc"
    if any(ext in w for ext in (".txt", ".vtt", ".srt")):
        return "text"
    if ".json" in w:
        return "json"
    return ""


def _kind_for_file(item: DriveFile) -> str:
    if item.mime_type == GDOC_MIME:
        return "gdoc"
    if item.mime_type == DOCX_MIME or item.name.lower().endswith(".docx"):
        return file_kind(item.name) or "docx"
    if item.name.lower().endswith((".txt", ".vtt", ".srt")):
        return "text"
    return file_kind(item.name)


def _docx_via_python(content: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def _docx_via_zip(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    parts: list[str] = []
    for node in root.findall(".//w:t", DOCX_NS):
        if node.text:
            parts.append(node.text)
        if node.tail:
            parts.append(node.tail)
    return "\n".join(p for p in parts if p).strip()


def docx_text(content: bytes) -> str:
    """Plain text from a call-transcriber .docx (python-docx, zip fallback)."""
    try:
        text = _docx_via_python(content)
        if text:
            return text
    except Exception:
        pass
    return _docx_via_zip(content)


def download_bytes(auth: DriveAuth, file_id: str) -> bytes:
    params = {"alt": "media", "supportsAllDrives": "true"}
    data = drive_get(auth, f"{DRIVE_FILES}/{file_id}", params)
    return data if isinstance(data, (bytes, bytearray)) else bytes(data)


def export_google_doc(auth: DriveAuth, file_id: str) -> str:
    params = {"mimeType": "text/plain"}
    data = drive_get(auth, f"{DRIVE_FILES}/{file_id}/export", params)
    if isinstance(data, (bytes, bytearray)):
        return data.decode("utf-8", errors="replace")
    if isinstance(data, str):
        return data
    return str(data)


def _occurred_at(meta: dict[str, str], item: DriveFile) -> datetime:
    """Title clock time is America/Chicago; store UTC."""
    date = meta.get("date") or item.folder_date
    time_s = meta.get("time") or "00:00:00"
    if date:
        parts = (time_s or "00:00:00").split(":")
        try:
            hh = int(parts[0])
            mm = int(parts[1]) if len(parts) > 1 else 0
            ss = int(parts[2]) if len(parts) > 2 else 0
            local = datetime.fromisoformat(date).replace(hour=hh, minute=mm, second=ss, tzinfo=CDT)
            return local.astimezone(timezone.utc)
        except ValueError:
            pass
    return item.modified_time or now_utc()


def _load_text(auth: DriveAuth, item: DriveFile, kind: str) -> str:
    if kind == "gdoc":
        return export_google_doc(auth, item.file_id)
    if kind in {"docx_transcript", "docx"}:
        return docx_text(download_bytes(auth, item.file_id))
    if kind == "text":
        raw = download_bytes(auth, item.file_id)
        return raw.decode("utf-8", errors="replace")
    return ""


def scan(
    settings: Settings,
    dates: Iterable[str] | None = None,
    *,
    backfill: bool = False,
) -> list[Engagement]:
    """Read Cube ACR via Drive API v3. Prefer .docx; export Google Docs as text."""
    auth = resolve_drive_auth(settings)
    try:
        candidates = list_transcript_candidates(
            settings, auth, dates=dates, backfill=backfill
        )
    except PermissionError as exc:
        if auth.mode == "oauth" and (settings.google_api_key or "").strip():
            logger.warning("cube drive oauth denied (%s); falling back to GOOGLE_API_KEY", exc)
            auth = DriveAuth(mode="api_key", headers={}, key=settings.google_api_key.strip())
            candidates = list_transcript_candidates(
                settings, auth, dates=dates, backfill=backfill
            )
        else:
            raise CubeAuthError(
                f"{exc}; grant {DRIVE_READONLY} or set GOOGLE_API_KEY"
            ) from exc

    parsed: list[tuple[DriveFile, str, dict[str, str], str]] = []
    seen_ids: set[str] = set()
    seen_md5: set[str] = set()
    for item in candidates:
        kind = _kind_for_file(item)
        if kind not in {"docx_transcript", "docx", "gdoc", "text"}:
            continue
        if item.file_id in seen_ids:
            continue
        if item.md5 and item.md5 in seen_md5:
            continue
        meta = parse_cube_title(item.name)
        name = meta.get("name") or ""
        phone = meta.get("phone") or ""
        if should_skip_cube_call(name=name, phone=phone, title=item.name):
            continue
        seen_ids.add(item.file_id)
        if item.md5:
            seen_md5.add(item.md5)
        parsed.append((item, kind, meta, cube_call_key(meta)))

    winners: dict[str, tuple[DriveFile, str, dict[str, str]]] = {}
    for item, kind, meta, key in parsed:
        prev = winners.get(key)
        if prev is None or _kind_rank(kind, item.name) < _kind_rank(prev[1], prev[0].name):
            winners[key] = (item, kind, meta)

    engagements: list[Engagement] = []
    for item, kind, meta in winners.values():
        name = meta.get("name") or ""
        phone = meta.get("phone") or ""
        try:
            text = _load_text(auth, item, kind)
        except Exception as exc:
            logger.warning("cube download %s: %s", item.file_id, exc)
            continue
        if looks_like_html(text) or len(text.strip()) < 20:
            continue
        first, _, last = name.partition(" ")
        engagements.append(
            Engagement(
                source="cube_acr",
                external_id=item.file_id,
                occurred_at=_occurred_at(meta, item),
                name=name,
                first_name=first,
                last_name=last.strip(),
                phone=phone,
                transcript=text[:20000],
                summary=text[:800],
                raw_subject=item.name,
                extra={
                    "folder_date": item.folder_date or meta.get("date") or "",
                    "transcript_kind": kind,
                    "md5": item.md5,
                    "call_time": meta.get("time") or "",
                    "call_key": cube_call_key(meta),
                    "drive_auth": auth.mode,
                },
            )
        )
    engagements.sort(key=lambda ev: ev.occurred_at or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
    return engagements
